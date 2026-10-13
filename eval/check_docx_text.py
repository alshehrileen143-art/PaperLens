"""How much visible DOCX text the extraction keeps, with the hidden-text filter off and on.

Ordinary documents are made here (no personal files), one kind of content each: links (DOI
references, links in the text, cross-references, a table of contents, an email), citations from a
reference manager (w:fldSimple), tracked insertions (w:ins), content controls (w:sdt) and a table.
The truth is every w:t of the body that is not deleted (w:del). "Lost" = visible characters missing
from extract_text's text. Files given with --files (e.g. the hand-made injection DOCX files) are
shown too, with the characters each path returns. Local only, no API.

--compare-with <git revision> runs the same with extract.py and hidden_text.py of that revision
(e.g. HEAD), for a before and after.

Run (PowerShell, from the project folder):
    .\\venv\\Scripts\\python.exe eval\\check_docx_text.py --compare-with HEAD --files samples\\injection
"""

import argparse
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent
CODE = os.environ.get("CHECK_DOCX_CODE")  # Set by --compare-with: a folder with older modules.
sys.path.insert(0, str(ROOT))
if CODE:
    sys.path.insert(0, CODE)

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402

import config  # noqa: E402
import extract  # noqa: E402

W = nsdecls("w")
LINK_RELATION = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
BODY = ("Vitamin D deficiency was more frequent in children with obesity than in controls, "
        "with a pooled relative risk of 1.41 across twenty studies.")


def run_xml(text: str, extra: str = "") -> str:
    return f'<w:r><w:rPr>{extra}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


def link(paragraph, text: str, url: str | None = None, anchor: str | None = None) -> None:
    target = f'w:anchor="{anchor}"' if anchor else ""
    element = parse_xml(f'<w:hyperlink {W} {target}>{run_xml(text, "<w:rStyle w:val=\'Hyperlink\'/>")}</w:hyperlink>')
    if url:
        rel = paragraph.part.relate_to(url, LINK_RELATION,
                                       is_external=True)
        element.set(qn("r:id"), rel)
    paragraph._p.append(element)


def wrap(paragraph, tag: str, text: str, attrs: str = "") -> None:
    inner = run_xml(text)
    if tag == "sdt":
        xml = f"<w:sdt {W}><w:sdtContent>{inner}</w:sdtContent></w:sdt>"
    elif tag == "fldSimple":
        xml = f'<w:fldSimple {W} w:instr=" ADDIN ZOTERO_ITEM CSL_CITATION ">{inner}</w:fldSimple>'
    else:
        xml = f'<w:{tag} {W} {attrs}>{inner}</w:{tag}>'
    paragraph._p.append(parse_xml(xml))


def documents() -> dict:
    made = {}

    def new(name):
        document = Document()
        made[name] = document
        return document

    d = new("references with DOI links")
    for n in range(1, 21):
        doi = f"https://doi.org/10.1016/j.jped.20{n:02d}.0{n:02d}"
        p = d.add_paragraph(f"{n}. Author A, Author B. A study of vitamin D in children. J Pediatr. 20{n:02d};{n}:1-9. ")
        link(p, doi, url=doi)
    d = new("links inside the text")
    for _ in range(10):
        p = d.add_paragraph("Data were taken from ")
        link(p, "the national survey portal", url="https://example.org/survey")
        p.add_run(" and from ")
        link(p, "PubMed", url="https://pubmed.ncbi.nlm.nih.gov/")
        p.add_run(". " + BODY)
    d = new("cross-references to figures and references")
    for _ in range(10):
        p = d.add_paragraph("As shown in ")
        link(p, "Figure 2", anchor="_Ref1")
        p.add_run(" and reported earlier ")
        link(p, "[12]", anchor="_Ref2")
        p.add_run(", " + BODY)
    d = new("table of contents")
    for title in ("1. Introduction", "2. Methods", "2.1 Search strategy", "3. Results", "4. Discussion",
                  "5. Conclusion"):
        link(d.add_paragraph(), f"{title}\t3", anchor="_Toc1")
    for _ in range(5):
        d.add_paragraph(BODY)
    d = new("email in the affiliations")
    p = d.add_paragraph("Corresponding author: ")
    link(p, "author@university.edu", url="mailto:author@university.edu")
    for _ in range(5):
        d.add_paragraph(BODY)
    d = new("citations from a reference manager (fldSimple)")
    for n in range(10):
        p = d.add_paragraph(BODY + " ")
        wrap(p, "fldSimple", f"(Smith et al., 20{n:02d})")
    d = new("tracked insertion (w:ins)")
    for _ in range(10):
        p = d.add_paragraph(BODY + " ")
        wrap(p, "ins", "This sentence was added in review.",
             'w:id="1" w:author="Reviewer" w:date="2026-01-01T00:00:00Z"')
    d = new("content control (w:sdt)")
    for _ in range(10):
        p = d.add_paragraph("Title: ")
        wrap(p, "sdt", "Vitamin D and obesity in children")
    d = new("table")
    table = d.add_table(rows=6, cols=3)
    for row in table.rows:
        for cell, value in zip(row.cells, ("Rajakumar et al.", "United States", "1.41 (1.26-1.59)")):
            cell.text = value
    d.add_paragraph(BODY)
    return made


def truth(path: Path) -> int:
    body = Document(path).element.body
    deleted = {id(t) for d in body.iter(qn("w:del")) for t in d.iter(qn("w:t"))}
    return sum(len(t.text or "") for t in body.iter(qn("w:t")) if id(t) not in deleted)


def chars(path: Path, hidden_filter: bool) -> int:
    with patch.object(config, "HIDDEN_TEXT_FILTER", hidden_filter):
        return len(extract.extract_text(path).text.replace("\n", ""))


def measure(files: list[Path]) -> list[dict]:
    rows = []
    with tempfile.TemporaryDirectory() as folder:
        for name, document in documents().items():
            path = Path(folder) / "doc.docx"
            document.save(path)
            expected = truth(path)
            rows.append({"name": name, "visible": expected, "off": expected - chars(path, False),
                         "on": expected - chars(path, True), "made": True})
    for path in files:
        rows.append({"name": path.name, "visible": truth(path), "off": chars(path, False), "on": chars(path, True),
                     "made": False})
    return rows


def at_revision(revision: str, files: list[Path]) -> list[dict]:
    with tempfile.TemporaryDirectory() as folder:
        for module in ("extract.py", "hidden_text.py"):
            source = subprocess.run(["git", "show", f"{revision}:{module}"], cwd=ROOT, capture_output=True,
                                    text=True, encoding="utf-8", check=True).stdout
            (Path(folder) / module).write_text(source, encoding="utf-8")
        done = subprocess.run([sys.executable, __file__, "--json", *(f"--files={p}" for p in files)],
                              capture_output=True, text=True, encoding="utf-8", check=True,
                              env={**os.environ, "CHECK_DOCX_CODE": folder, "PYTHONIOENCODING": "utf-8"})
    return json.loads(done.stdout.strip().splitlines()[-1])


def table(rows: list[dict], before: list[dict] | None) -> str:
    def cell(n, key):
        now = rows[n][key]
        return f"{before[n][key]} → {now}" if before and before[n][key] != now else str(now)
    lines = ["| Document | Visible characters | Lost, filter off | Lost, filter on |", "|---|---|---|---|"]
    lines += [f"| {r['name']} | {r['visible']} | {cell(n, 'off')} | {cell(n, 'on')} |"
              for n, r in enumerate(rows) if r["made"]]
    given = [n for n, r in enumerate(rows) if not r["made"]]
    if given:
        lines += ["", "| File | w:t characters (not deleted) | Text, filter off | Text, filter on |",
                  "|---|---|---|---|"]
        lines += [f"| {rows[n]['name']} | {rows[n]['visible']} | {cell(n, 'off')} | {cell(n, 'on')} |" for n in given]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--files", type=Path, action="append", default=[], help="DOCX files or folders.")
    parser.add_argument("--compare-with", help="A git revision whose extract.py and hidden_text.py are the 'before'.")
    parser.add_argument("--json", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    files = sorted(p for item in args.files for p in (item.glob("*.docx") if item.is_dir() else [item]))
    rows = measure(files)
    if args.json:
        print(json.dumps(rows))
        return
    before = at_revision(args.compare_with, files) if args.compare_with else None
    print(table(rows, before))


if __name__ == "__main__":
    main()
