"""DOCX tables in extraction: the two papers rebuilt as DOCX with their real tables as Word tables,
measured through extract_text (filter on): characters, tokens, mode (whole paper or RAG), RAG
chunks, full-context sections, the sections detected, quotes of table rows (cells joined by spaces)
that quote_in_source finds, and table rows cut by a RAG chunk boundary. Local only, no API.

The DOCX files are made from the PDFs: text outside the tables as paragraphs (headings as paragraphs
of their own), each table as a Word table where its lines were, its cells split where words are
more than a few points apart, its caption a paragraph before it. --compare-with <git revision> runs
the same with extract.py and hidden_text.py of that revision (e.g. HEAD), for a before and after.

Run (PowerShell, from the project folder):
    .\\venv\\Scripts\\python.exe eval\\check_docx_tables.py --compare-with HEAD --migraine <path to the migraine PDF>
"""

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent
CODE = os.environ.get("CHECK_DOCX_CODE")  # Set by --compare-with: a folder with older modules.
sys.path.insert(0, str(ROOT))
if CODE:
    sys.path.insert(0, CODE)

import pymupdf  # noqa: E402
from docx import Document  # noqa: E402
from pypdf import PdfReader  # noqa: E402

import config  # noqa: E402
import extract  # noqa: E402
import rag  # noqa: E402
from llm import count_tokens  # noqa: E402
from quote_match import quote_in_source  # noqa: E402
from sections import _heading_name, detect_sections  # noqa: E402
from ui_text import pdf_paragraphs  # noqa: E402

VITAMIN_D = ROOT / "samples" / "1-s2.0-S0021755720302072-main.pdf"
# (page index, table area in display coordinates, gap between cells in points), found by hand.
VITAMIN_D_TABLES = [(2, pymupdf.Rect(0, 60, 800, 560), 6)]  # Page 3 is rotated (landscape).
MIGRAINE_TABLES = [(2, pymupdf.Rect(305, 70, 560, 506), 12), (4, pymupdf.Rect(40, 70, 560, 586), 12)]


def table_rows(page, clip, gap):
    """Rows of words in display coordinates, each split into cells at gaps wider than `gap` points."""
    words = []
    for w in page.get_text("words"):
        r = pymupdf.Rect(w[:4]) * page.rotation_matrix
        if r.intersects(clip):
            words.append((r, w[4]))
    words.sort(key=lambda t: ((t[0].y0 + t[0].y1) / 2, t[0].x0))
    rows = []
    for r, text in words:
        middle = (r.y0 + r.y1) / 2
        if rows and abs(rows[-1]["y"] - middle) < 4:
            rows[-1]["words"].append((r, text))
        else:
            rows.append({"y": middle, "words": [(r, text)]})
    out = []
    for row in rows:
        ws = sorted(row["words"], key=lambda t: t[0].x0)
        cells, current = [], [ws[0]]
        for a, b in zip(ws, ws[1:]):
            if b[0].x0 - a[0].x1 > gap:
                cells.append(current)
                current = []
            current.append(b)
        cells.append(current)
        out.append([" ".join(t for _, t in cell) for cell in cells])
    return out


def build_docx(pdf: Path, tables: list, out: Path) -> list[list[list[str]]]:
    """The paper as DOCX: text outside the tables as paragraphs, each table as a Word table where
    its lines were. Returns the tables' rows."""
    document = Document()
    all_tables = []
    reader = PdfReader(pdf)
    with pymupdf.open(pdf) as source:
        for index, page_reader in enumerate(reader.pages):
            text = page_reader.extract_text() or ""
            here = [(clip, gap) for page_index, clip, gap in tables if page_index == index]
            rows_here = [table_rows(source[index], clip, gap) for clip, gap in here]
            table_text = "".join("".join(cell for row in rows for cell in row) for rows in rows_here)
            squashed = re.sub(r"\s", "", table_text)
            kept, placed = [], False
            for line in text.split("\n"):
                key = re.sub(r"\s", "", line)
                if rows_here and len(key) >= 3 and key in squashed:
                    if not placed:
                        kept.append(None)  # Where the tables go.
                        placed = True
                    continue
                kept.append(line)
            segments, current = [], []
            for line in kept:
                if line is None:
                    segments.append(("text", "\n".join(current)))
                    segments.append(("tables", None))
                    current = []
                else:
                    current.append(line)
            segments.append(("text", "\n".join(current)))
            for kind, value in segments:
                if kind == "text":
                    # Headings stay paragraphs of their own, as in a real DOCX.
                    block = []
                    for line in value.split("\n") + [None]:
                        if line is None or _heading_name(line):
                            for paragraph in pdf_paragraphs("\n".join(block)):
                                document.add_paragraph(paragraph)
                            block = []
                            if line is not None:
                                document.add_paragraph(line.strip())
                        else:
                            block.append(line)
                else:
                    for rows in rows_here:
                        # The caption ("Table 1: ..." and its continuation) is a paragraph before the table.
                        caption = []
                        def is_caption(row):
                            return bool(re.match(r"(?i)table\s*\d", row[0])) or (
                                bool(caption) and len(row) == 1 and not re.search(r"\d", row[0]))
                        while rows and is_caption(rows[0]):
                            caption.append(" ".join(rows.pop(0)))
                        if caption:
                            document.add_paragraph(" ".join(caption))
                        width = max(len(row) for row in rows)
                        table = document.add_table(rows=len(rows), cols=width)
                        for r, row in enumerate(rows):
                            for c, cell in enumerate(row):
                                table.cell(r, c).text = cell
                        all_tables.append(rows)
    document.save(out)
    return all_tables


def measure(path: Path, tables: list) -> dict:
    with patch.object(config, "HIDDEN_TEXT_FILTER", True):
        text = extract.extract_text(path).text
    sections = detect_sections(text)
    chunks = rag.chunk_text(text, None, label_sections=sections)
    tokens = count_tokens(text)
    parts = rag.full_context({"text": text, "pages": None, "token_count": tokens})
    quotes = [" ".join(row) for rows in tables for row in rows if len(row) >= 2][:40]
    row_lines = [line for line in text.split("\n") if "\t" in line]
    return {"chars": len(text), "tokens": tokens,
            "mode": "full" if tokens <= config.FULL_CONTEXT_MAX_TOKENS else "rag", "chunks": len(chunks),
            "full_parts": len(parts), "sections": ", ".join(name for _, name in sections),
            "quotes": f"{sum(quote_in_source(q, text) for q in quotes)}/{len(quotes)}",
            "rows_cut": f"{sum(1 for line in row_lines if not any(line in c.text for c in chunks))}/{len(row_lines)}"}


def measure_all(papers: dict) -> dict:
    results = {}
    with tempfile.TemporaryDirectory() as folder:
        for name, (pdf, tables) in papers.items():
            path = Path(folder) / f"{name}.docx"
            rows = build_docx(pdf, tables, path)
            results[name] = {"tables": len(rows), "table_chars": sum(len(c) for r in rows for x in r for c in x),
                             **measure(path, rows)}
    return results


def at_revision(revision: str, arguments: list[str]) -> dict:
    with tempfile.TemporaryDirectory() as folder:
        for module in ("extract.py", "hidden_text.py"):
            source = subprocess.run(["git", "show", f"{revision}:{module}"], cwd=ROOT, capture_output=True,
                                    text=True, encoding="utf-8", check=True).stdout
            (Path(folder) / module).write_text(source, encoding="utf-8")
        done = subprocess.run([sys.executable, __file__, "--json", *arguments], capture_output=True, text=True,
                              encoding="utf-8", check=True,
                              env={**os.environ, "CHECK_DOCX_CODE": folder, "PYTHONIOENCODING": "utf-8"})
    return json.loads(done.stdout.strip().splitlines()[-1])


def table(now: dict, before: dict | None) -> str:
    keys = [("chars", "Characters"), ("tokens", "Tokens"), ("mode", "Mode"), ("chunks", "RAG chunks"),
            ("full_parts", "Full-context sections"), ("quotes", "Table-row quotes found"),
            ("rows_cut", "Table rows cut by a chunk")]
    lines = []
    for name, row in now.items():
        lines += [f"### {name} ({row['tables']} table(s), {row['table_chars']} characters in them)", "",
                  "| Measure | " + ("Before | After |" if before else "Value |"),
                  "|---|---|" + ("---|" if before else "")]
        for key, label in keys:
            lines.append(f"| {label} | " + (f"{before[name][key]} | {row[key]} |" if before else f"{row[key]} |"))
        lines += ["", f"Sections now: {row['sections']}"]
        if before and before[name]["sections"] != row["sections"]:
            lines.append(f"Sections before: {before[name]['sections']}")
        lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--migraine", type=Path, help="The migraine paper's PDF (not in the repository).")
    parser.add_argument("--compare-with", help="A git revision whose extract.py and hidden_text.py are the 'before'.")
    parser.add_argument("--json", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    papers = {"vitamin D": (VITAMIN_D, VITAMIN_D_TABLES)}
    if args.migraine:
        papers["migraine"] = (args.migraine, MIGRAINE_TABLES)
    now = measure_all(papers)
    if args.json:
        print(json.dumps(now))
        return
    forwarded = [f"--migraine={args.migraine}"] if args.migraine else []
    before = at_revision(args.compare_with, forwarded) if args.compare_with else None
    print(table(now, before))


if __name__ == "__main__":
    main()
